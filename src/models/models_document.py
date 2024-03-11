import os

import pathlib
import docx
import json
from docxtpl import DocxTemplate

from .models_exceptions import *


class Document:
    """
        class Document, 2024-02-27
    """
    document: docx.document.Document | None = None

    def __init__(self, **kwargs: dict):
        self.kwargs: dict = kwargs
        self.set_attrs()
        self.open_template()

    def set_attrs(self):
        for k, v in self.kwargs.items():
            setattr(self, k, v)
        del self.kwargs

    def open_template(self):
        try:
            self.document = docx.Document(self.template_filepath)
            return
        except Exception:
            self.document = None
            message = f'\n{self.template_filepath}_NOT_FOUND'
            raise TemplateNotFoundError(message=message)

    def populate_placeholder(self, placeholders_filepath):
        try:
            with open(placeholders_filepath, 'r') as file:
                placeholders = json.load(file)
        except Exception:
            return

        template = DocxTemplate(self.template_filepath)
        template.render(placeholders)
        template.save(self.output_filepath)
        os.remove(placeholders_filepath)
